import argparse
import hashlib
import logging
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

from pptx import Presentation
from pptx.shapes.group import GroupShape


HARDCODED_INPUT_DIR = "input"
HARDCODED_OUTPUT_FILE = "all_powerpoint_text.txt"
HARDCODED_LOG_FILE = "powerpoint_text_extractor.log"
HARDCODED_RECURSIVE = False

SUPPORTED_EXTENSIONS = {".ppt", ".pptx"}


@dataclass
class ExtractedSlide:
    slide_number: int
    text_blocks: List[str]


@dataclass
class ExtractedPresentation:
    source_path: Path
    converted_path: Optional[Path]
    slides: List[ExtractedSlide]
    error: str = ""


def setup_logger(log_path: Path) -> logging.Logger:
    log_path.parent.mkdir(parents=True, exist_ok=True)

    logger = logging.getLogger("powerpoint_text_extractor")
    logger.setLevel(logging.INFO)

    if not logger.handlers:
        formatter = logging.Formatter(
            "%(asctime)s [%(levelname)s] %(message)s",
            datefmt="%Y-%m-%d %H:%M:%S",
        )

        file_handler = logging.FileHandler(log_path, encoding="utf-8")
        file_handler.setFormatter(formatter)
        logger.addHandler(file_handler)

        console_handler = logging.StreamHandler()
        console_handler.setFormatter(formatter)
        logger.addHandler(console_handler)

    return logger


def short_hash(text: str, length: int = 8) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:length]


def find_powerpoint_files(input_dir: Path, recursive: bool) -> List[Path]:
    pattern = "**/*" if recursive else "*"

    files = [
        path
        for path in input_dir.glob(pattern)
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
        and not path.name.startswith("~$")
    ]

    return sorted(files)


def libreoffice_available() -> bool:
    return shutil.which("libreoffice") is not None or shutil.which("soffice") is not None


def get_libreoffice_command() -> str:
    command = shutil.which("libreoffice") or shutil.which("soffice")

    if not command:
        raise RuntimeError(
            "LibreOffice was not found. Install LibreOffice to process legacy .ppt files."
        )

    return command


def convert_ppt_to_pptx(
    ppt_path: Path,
    conversion_dir: Path,
    logger: logging.Logger,
) -> Path:
    command = get_libreoffice_command()

    logger.info("Converting .ppt to .pptx: %s", ppt_path)

    profile_dir = conversion_dir / "lo_profile"

    result = subprocess.run(
        [
            command,
            "--headless",
            f"-env:UserInstallation={profile_dir.resolve().as_uri()}",
            "--convert-to",
            "pptx",
            "--outdir",
            str(conversion_dir),
            str(ppt_path),
        ],
        capture_output=True,
        text=True,
        timeout=120,
    )

    if result.returncode != 0:
        raise RuntimeError(
            f"LibreOffice conversion failed for {ppt_path}\n"
            f"STDOUT: {result.stdout}\n"
            f"STDERR: {result.stderr}"
        )

    expected_path = conversion_dir / f"{ppt_path.stem}.pptx"

    if expected_path.exists():
        return expected_path

    converted_candidates = list(conversion_dir.glob("*.pptx"))

    if not converted_candidates:
        raise RuntimeError(f"No converted .pptx file found for {ppt_path}")

    return converted_candidates[0]


def extract_text_from_shape(shape) -> List[str]:
    text_blocks: List[str] = []

    if hasattr(shape, "text") and shape.text:
        text = shape.text.strip()

        if text:
            text_blocks.append(text)

    if shape.has_table:
        table_lines = []

        for row in shape.table.rows:
            cells = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    cells.append(cell_text)

            if cells:
                table_lines.append(" | ".join(cells))

        if table_lines:
            text_blocks.append("\n".join(table_lines))

    if isinstance(shape, GroupShape):
        for grouped_shape in shape.shapes:
            text_blocks.extend(extract_text_from_shape(grouped_shape))

    return text_blocks


def extract_notes_text(slide) -> List[str]:
    notes_blocks: List[str] = []

    if not slide.has_notes_slide:
        return notes_blocks

    notes_slide = slide.notes_slide

    for paragraph in notes_slide.notes_text_frame.paragraphs:
        paragraph_text = paragraph.text.strip()

        if paragraph_text:
            notes_blocks.append(paragraph_text)

    return notes_blocks


def extract_text_from_pptx(pptx_path: Path, logger: logging.Logger) -> List[ExtractedSlide]:
    presentation = Presentation(str(pptx_path))

    extracted_slides: List[ExtractedSlide] = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        text_blocks: List[str] = []

        for shape in slide.shapes:
            try:
                text_blocks.extend(extract_text_from_shape(shape))
            except Exception as exc:
                logger.warning(
                    "Skipping unreadable shape on slide %d of %s: %s",
                    slide_index,
                    pptx_path,
                    exc,
                )

        try:
            notes_blocks = extract_notes_text(slide)
        except Exception as exc:
            notes_blocks = []
            logger.warning(
                "Failed to extract speaker notes on slide %d of %s: %s",
                slide_index,
                pptx_path,
                exc,
            )

        if notes_blocks:
            text_blocks.append("[Speaker Notes]\n" + "\n".join(notes_blocks))

        cleaned_blocks = []

        for block in text_blocks:
            cleaned = "\n".join(
                line.strip()
                for line in block.splitlines()
                if line.strip()
            )

            if cleaned:
                cleaned_blocks.append(cleaned)

        extracted_slides.append(
            ExtractedSlide(
                slide_number=slide_index,
                text_blocks=cleaned_blocks,
            )
        )

    return extracted_slides


def extract_presentation(
    source_path: Path,
    temp_dir: Path,
    logger: logging.Logger,
) -> ExtractedPresentation:
    converted_path: Optional[Path] = None

    try:
        if source_path.suffix.lower() == ".pptx":
            slides = extract_text_from_pptx(source_path, logger)

        elif source_path.suffix.lower() == ".ppt":
            conversion_dir = temp_dir / f"converted_{source_path.stem}_{short_hash(str(source_path))}"
            conversion_dir.mkdir(parents=True, exist_ok=True)

            converted_path = convert_ppt_to_pptx(
                ppt_path=source_path,
                conversion_dir=conversion_dir,
                logger=logger,
            )

            slides = extract_text_from_pptx(converted_path, logger)

        else:
            raise ValueError(f"Unsupported file type: {source_path.suffix}")

        return ExtractedPresentation(
            source_path=source_path,
            converted_path=converted_path,
            slides=slides,
        )

    except Exception as exc:
        logger.error("Failed to extract text from %s: %s", source_path, exc)

        return ExtractedPresentation(
            source_path=source_path,
            converted_path=converted_path,
            slides=[],
            error=str(exc),
        )


def write_combined_text_file(
    output_path: Path,
    presentations: List[ExtractedPresentation],
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    total_files = len(presentations)
    successful_files = sum(1 for item in presentations if not item.error)
    failed_files = sum(1 for item in presentations if item.error)
    total_slides = sum(len(item.slides) for item in presentations)

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("POWERPOINT TEXT EXTRACTION\n")
        f.write("=" * 80 + "\n")
        f.write(f"Total files processed: {total_files}\n")
        f.write(f"Successful files: {successful_files}\n")
        f.write(f"Failed files: {failed_files}\n")
        f.write(f"Total slides extracted: {total_slides}\n")
        f.write("=" * 80 + "\n\n")

        for presentation in presentations:
            f.write("\n")
            f.write("#" * 80 + "\n")
            f.write(f"FILE: {presentation.source_path}\n")

            if presentation.converted_path:
                f.write(f"CONVERTED_FROM_PPT_TO: {presentation.converted_path}\n")

            if presentation.error:
                f.write(f"ERROR: {presentation.error}\n")
                f.write("#" * 80 + "\n\n")
                continue

            f.write(f"SLIDES: {len(presentation.slides)}\n")
            f.write("#" * 80 + "\n\n")

            for slide in presentation.slides:
                f.write("-" * 80 + "\n")
                f.write(f"SLIDE {slide.slide_number}\n")
                f.write("-" * 80 + "\n")

                if not slide.text_blocks:
                    f.write("[No extractable text found]\n\n")
                    continue

                for block in slide.text_blocks:
                    f.write(block)
                    f.write("\n\n")


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Batch extract text from .ppt and .pptx files into one combined .txt file."
    )

    parser.add_argument(
        "--input-dir",
        default=None,
        help=(
            "Optional directory containing .ppt and .pptx files. "
            "If omitted, HARDCODED_INPUT_DIR is used."
        ),
    )

    parser.add_argument(
        "--output-file",
        default=None,
        help=(
            "Optional path to the combined output .txt file. "
            "If omitted, HARDCODED_OUTPUT_FILE is used."
        ),
    )

    parser.add_argument(
        "--log-file",
        default=None,
        help=(
            "Optional path to the log file. "
            "If omitted, HARDCODED_LOG_FILE is used."
        ),
    )

    parser.add_argument(
        "--recursive",
        action="store_true",
        help=(
            "Recursively scan input directory. "
            "If omitted, HARDCODED_RECURSIVE is used."
        ),
    )

    return parser


def main() -> None:
    parser = build_arg_parser()
    args = parser.parse_args()

    input_dir = Path(args.input_dir or HARDCODED_INPUT_DIR)
    output_file = Path(args.output_file or HARDCODED_OUTPUT_FILE)
    log_file = Path(args.log_file or HARDCODED_LOG_FILE)

    recursive = args.recursive or HARDCODED_RECURSIVE

    logger = setup_logger(log_file)

    logger.info("=== PowerPoint text extraction started ===")
    logger.info("Input directory: %s", input_dir.resolve())
    logger.info("Output file: %s", output_file.resolve())
    logger.info("Log file: %s", log_file.resolve())
    logger.info("Recursive: %s", recursive)

    if not input_dir.exists():
        raise SystemExit(f"Input directory does not exist: {input_dir}")

    files = find_powerpoint_files(input_dir=input_dir, recursive=recursive)

    if not files:
        raise SystemExit(f"No .ppt or .pptx files found in: {input_dir}")

    logger.info("Found %d PowerPoint files in %s", len(files), input_dir)

    has_ppt_files = any(path.suffix.lower() == ".ppt" for path in files)

    if has_ppt_files and not libreoffice_available():
        logger.warning(
            "Legacy .ppt files were found, but LibreOffice is not available. "
            ".ppt files will fail unless LibreOffice is installed."
        )

    presentations: List[ExtractedPresentation] = []

    with tempfile.TemporaryDirectory() as temp_root:
        temp_dir = Path(temp_root)

        for index, file_path in enumerate(files, start=1):
            print(f"[{index}/{len(files)}] Extracting: {file_path}")
            logger.info("Extracting file %d/%d: %s", index, len(files), file_path)

            extracted = extract_presentation(
                source_path=file_path,
                temp_dir=temp_dir,
                logger=logger,
            )

            presentations.append(extracted)

    write_combined_text_file(
        output_path=output_file,
        presentations=presentations,
    )

    successful = sum(1 for item in presentations if not item.error)
    failed = sum(1 for item in presentations if item.error)

    logger.info("Extraction complete. Output written to: %s", output_file)
    logger.info(
        "Summary: total=%d successful=%d failed=%d",
        len(presentations),
        successful,
        failed,
    )
    logger.info("=== PowerPoint text extraction finished ===\n")

    print(f"Text written to: {output_file.resolve()}")
    print(f"Summary: total={len(presentations)}, successful={successful}, failed={failed}")

    if failed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()